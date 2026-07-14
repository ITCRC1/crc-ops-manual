"""Laboratorio de intake: triage por similitud (posible duplicado vs nuevo) + extracción
de texto de archivos .txt/.docx."""

import io

import pytest

pytestmark = pytest.mark.asyncio(loop_scope="session")

ADMIN_EMAIL = "Finance@thecostaricacollection.com"
ADMIN_PASS = "CambiarCRC2026!"


async def _headers(client):
    r = await client.post("/api/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASS})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


async def test_requiere_editor(client):
    assert (await client.post("/api/intake/analizar", json={"nombre": "x", "texto": ""})).status_code == 403


async def test_analizar_duplicado_vs_nuevo(client):
    from datetime import UTC, datetime

    from sqlalchemy import select

    from app.core.db import AsyncSessionLocal
    from app.models.sop import SopMaster

    h = await _headers(client)
    nombre = "Procedimiento Único de Prueba Intake XYZ"
    texto = "reconciliación de caja chica arqueos diarios y liquidación de vales"
    sop = (await client.post(
        "/api/sops", json={"area_codigo": "GOV", "nombre": nombre, "proposito": texto}, headers=h
    )).json()
    try:
        # mismo nombre + contenido → posible duplicado, con ese SOP arriba
        r = await client.post("/api/intake/analizar", json={"nombre": nombre, "texto": texto}, headers=h)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["verdict"] == "posible_duplicado"
        assert any(s["id"] == sop["id"] and s["score"] > 0.8 for s in d["similares"])
        # algo totalmente distinto → nuevo
        r2 = await client.post(
            "/api/intake/analizar",
            json={
                "nombre": "Zzqwx evacuacion volcanica submarina 9910",
                "texto": "lava peces astronautas kdlfj",
            },
            headers=h,
        )
        assert r2.json()["verdict"] == "nuevo", r2.json()
    finally:
        async with AsyncSessionLocal() as db:
            s = (await db.execute(select(SopMaster).where(SopMaster.id == sop["id"]))).scalar_one_or_none()
            if s:
                s.eliminado_en = datetime.now(UTC)
                await db.commit()


async def test_extraer_txt(client):
    h = await _headers(client)
    files = {"archivo": ("proc.txt", b"Titulo del Procedimiento\nPaso 1\nPaso 2", "text/plain")}
    r = await client.post("/api/intake/extraer", files=files, headers=h)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["nombre"] == "Titulo del Procedimiento"
    assert "Paso 1" in d["texto"]


async def test_extraer_docx(client):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Manual de Compras")
    doc.add_paragraph("Contenido del manual de compras")
    buf = io.BytesIO()
    doc.save(buf)
    h = await _headers(client)
    files = {
        "archivo": (
            "manual.docx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    r = await client.post("/api/intake/extraer", files=files, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["nombre"] == "Manual de Compras"
    assert "Contenido" in r.json()["texto"]


async def test_extraer_pdf(client):
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, "Manual de Seguridad")
    c.drawString(100, 730, "Contenido de seguridad ocupacional")
    c.save()
    h = await _headers(client)
    files = {"archivo": ("manual.pdf", buf.getvalue(), "application/pdf")}
    r = await client.post("/api/intake/extraer", files=files, headers=h)
    assert r.status_code == 200, r.text
    assert "Seguridad" in r.json()["texto"]


async def test_extraer_formato_no_soportado(client):
    h = await _headers(client)
    files = {"archivo": ("hoja.xlsx", b"PK\x03\x04 fake", "application/vnd.ms-excel")}
    r = await client.post("/api/intake/extraer", files=files, headers=h)
    assert r.status_code == 415
